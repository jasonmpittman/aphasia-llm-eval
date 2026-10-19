from __future__ import annotations

import csv
import io
import json
import re
import zipfile

from openpyxl import load_workbook
from sqlalchemy import select

from aphasia_eval.db import Result, Run

HEADERS = {"Host": "localhost"}


def _completed_run(client) -> str:
    response = client.post(
        "/api/runs",
        json={
            "models": ["good-model", "alignment-model"],
            "tasks": ["ciu_scoring", "pos_breakdown"],
            "settings": {
                "temperature": 0,
                "max_tokens": 100,
                "repetitions": 1,
                "max_items_per_task": 2,
            },
            "label": "M5 synthetic export",
        },
        headers=HEADERS,
    )
    assert response.status_code == 202
    return response.get_json()["run_id"]


def test_json_export_contains_complete_versioned_run_without_secret(client):
    run_id = _completed_run(client)

    response = client.get(f"/api/runs/{run_id}/export?format=json", headers=HEADERS)

    assert response.status_code == 200
    assert response.mimetype == "application/json"
    assert response.headers["Cache-Control"] == "no-store"
    assert f"aphasia-eval-{run_id}.json" in response.headers["Content-Disposition"]
    assert b"unit-test-placeholder" not in response.data
    document = json.loads(response.data)
    assert document["schema_version"] == 3
    assert document["run"]["id"] == run_id
    assert document["run"]["label"] == "M5 synthetic export"
    assert document["run"]["status"] == "completed"
    assert len(document["models"]) == 2
    assert {task["task_id"] for task in document["tasks"]} == {
        "ciu_scoring",
        "pos_breakdown",
    }
    assert len(document["items"]) == 8
    assert document["aggregates"]
    assert all(len(task["gold_file_sha256"]) == 64 for task in document["tasks"])
    assert "not a clinical tool" in document["caveats"]["clinical_use"]
    assert "estimates" in document["caveats"]["cost"]


def test_csv_export_is_zip_with_item_and_aggregate_files(client):
    run_id = _completed_run(client)

    response = client.get(f"/api/runs/{run_id}/export?format=csv", headers=HEADERS)

    assert response.status_code == 200
    assert response.mimetype == "application/zip"
    assert f"aphasia-eval-{run_id}.csv.zip" in response.headers["Content-Disposition"]
    assert b"unit-test-placeholder" not in response.data
    with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
        assert set(archive.namelist()) == {"items.csv", "aggregates.csv"}
        items = list(csv.DictReader(io.TextIOWrapper(archive.open("items.csv"), "utf-8")))
        aggregates = list(csv.DictReader(io.TextIOWrapper(archive.open("aggregates.csv"), "utf-8")))
    assert len(items) == 8
    assert {row["task_id"] for row in items} == {"ciu_scoring", "pos_breakdown"}
    assert "comparison_json" in items[0]
    assert "estimated_cost_usd" in items[0]
    assert len(aggregates) == 4
    assert "estimated_cost_usd" in aggregates[0]
    assert "p95_latency_ms" in aggregates[0]
    assert any(column.startswith("headline") for column in aggregates[0])


def test_xlsx_export_has_required_sheets_and_rows(client):
    run_id = _completed_run(client)

    response = client.get(f"/api/runs/{run_id}/export?format=xlsx", headers=HEADERS)

    assert response.status_code == 200
    assert response.mimetype == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert b"unit-test-placeholder" not in response.data
    workbook = load_workbook(io.BytesIO(response.data), read_only=True, data_only=True)
    assert workbook.sheetnames == [
        "Summary",
        "Efficiency",
        "Agreement",
        "Confusion",
        "Items",
        "Run Info",
    ]
    assert workbook["Items"].max_row == 9
    assert workbook["Efficiency"].max_row == 5
    assert workbook["Agreement"].max_row > 1
    assert workbook["Confusion"].max_row > 1
    item_headers = [cell.value for cell in next(workbook["Items"].iter_rows())]
    assert "tokens_estimated" in item_headers
    assert "estimated_cost_usd" in item_headers
    summary_values = [cell.value for row in workbook["Summary"].iter_rows() for cell in row]
    assert "Research evaluation only; not a clinical tool." in summary_values
    assert "All cost values are estimated." in summary_values
    agreement_headers = [cell.value for cell in next(workbook["Agreement"].iter_rows())]
    assert agreement_headers[:8] == [
        "Model",
        "Task",
        "Condition",
        "Label",
        "Precision",
        "Recall",
        "F1",
        "Support",
    ]
    workbook.close()


def test_pdf_export_has_multiple_pages_and_required_sections(client):
    run_id = _completed_run(client)

    response = client.get(f"/api/runs/{run_id}/export?format=pdf", headers=HEADERS)

    assert response.status_code == 200
    assert response.mimetype == "application/pdf"
    assert response.data.startswith(b"%PDF-")
    assert b"unit-test-placeholder" not in response.data
    assert len(re.findall(rb"/Type\s*/Page\b", response.data)) >= 3


def test_export_validation_and_missing_run(client):
    missing_format = client.get("/api/runs/missing/export", headers=HEADERS)
    unsupported = client.get("/api/runs/missing/export?format=xml", headers=HEADERS)
    missing_run = client.get("/api/runs/missing/export?format=json", headers=HEADERS)

    assert missing_format.status_code == 400
    assert missing_format.get_json()["error"]["code"] == "invalid_format"
    assert unsupported.status_code == 400
    assert unsupported.get_json()["error"]["code"] == "invalid_format"
    assert missing_run.status_code == 404
    assert missing_run.get_json()["error"]["code"] == "not_found"


def test_ablation_exports_include_every_condition(client):
    response = client.post(
        "/api/runs",
        json={
            "models": ["good-model"],
            "tasks": ["ciu_scoring"],
            "settings": {
                "max_tokens": 512,
                "repetitions": 1,
                "max_items_per_task": 1,
                "ablation": {"parameter": "max_tokens", "values": [128]},
            },
        },
        headers=HEADERS,
    )
    run_id = response.get_json()["run_id"]

    json_response = client.get(f"/api/runs/{run_id}/export?format=json", headers=HEADERS)
    document = json.loads(json_response.data)
    assert {item["condition_id"] for item in document["items"]} == {
        "baseline",
        "max_tokens_128",
    }
    assert {row["condition_id"] for row in document["aggregates"]} == {
        "baseline",
        "max_tokens_128",
    }

    csv_response = client.get(f"/api/runs/{run_id}/export?format=csv", headers=HEADERS)
    with zipfile.ZipFile(io.BytesIO(csv_response.data)) as archive:
        items = list(csv.DictReader(io.TextIOWrapper(archive.open("items.csv"), "utf-8")))
        aggregates = list(csv.DictReader(io.TextIOWrapper(archive.open("aggregates.csv"), "utf-8")))
    assert {row["condition_id"] for row in items} == {"baseline", "max_tokens_128"}
    assert {row["condition_id"] for row in aggregates} == {
        "baseline",
        "max_tokens_128",
    }

    xlsx_response = client.get(f"/api/runs/{run_id}/export?format=xlsx", headers=HEADERS)
    workbook = load_workbook(io.BytesIO(xlsx_response.data), read_only=True, data_only=True)
    assert "Ablation" in workbook.sheetnames
    assert workbook["Ablation"].max_row == 3
    workbook.close()

    pdf_response = client.get(f"/api/runs/{run_id}/export?format=pdf", headers=HEADERS)
    assert pdf_response.status_code == 200
    assert pdf_response.data.startswith(b"%PDF-")


def test_tabular_exports_neutralize_spreadsheet_formulas(app, client):
    run_id = _completed_run(client)
    with app.extensions["aphasia_eval_database"].session() as session:
        run = session.get(Run, run_id)
        result = session.scalar(select(Result).where(Result.run_id == run_id).limit(1))
        assert run is not None and result is not None
        run.label = "=malicious-label"
        result.raw_output = '=HYPERLINK("https://attacker.invalid")'

    csv_response = client.get(f"/api/runs/{run_id}/export?format=csv", headers=HEADERS)
    with zipfile.ZipFile(io.BytesIO(csv_response.data)) as archive:
        items = list(csv.DictReader(io.TextIOWrapper(archive.open("items.csv"), "utf-8")))
    assert any(row["raw_output"].startswith("'=") for row in items)

    xlsx_response = client.get(f"/api/runs/{run_id}/export?format=xlsx", headers=HEADERS)
    workbook = load_workbook(io.BytesIO(xlsx_response.data), read_only=True, data_only=False)
    assert workbook["Summary"]["B2"].value == "'=malicious-label"
    raw_output_column = next(
        cell.column for cell in next(workbook["Items"].iter_rows()) if cell.value == "raw_output"
    )
    assert any(
        str(row[raw_output_column - 1].value).startswith("'=")
        for row in list(workbook["Items"].iter_rows())[1:]
    )
    workbook.close()
