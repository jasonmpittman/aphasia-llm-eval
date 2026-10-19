"""Build safe, reproducible run exports in the four supported formats."""

from __future__ import annotations

import csv
import json
import math
import re
import zipfile
from dataclasses import dataclass
from datetime import UTC, datetime
from io import BytesIO, StringIO
from typing import Any
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.shapes import Drawing
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import select
from sqlalchemy.orm import Session

from aphasia_eval.db import Aggregate, Result, Run, RunModel, RunTask
from aphasia_eval.services.ablation import result_experiment

CSV_ITEM_COLUMNS = (
    "result_id",
    "run_id",
    "model_id",
    "task_id",
    "item_id",
    "condition_id",
    "condition_json",
    "repetition",
    "input_json",
    "gold_json",
    "prompt_rendered",
    "raw_output",
    "parsed_output_json",
    "derived_json",
    "comparison_json",
    "matched",
    "input_tokens",
    "output_tokens",
    "reasoning_tokens",
    "tokens_estimated",
    "estimated_cost_usd",
    "cost_source",
    "latency_ms",
    "total_elapsed_ms",
    "attempts",
    "status",
    "error_message",
    "created_at",
)
CSV_AGGREGATE_ID_COLUMNS = ("run_id", "model_id", "task_id", "condition_id")
FORMULA_PREFIXES = ("=", "+", "-", "@")
PDF_BLUE = colors.HexColor("#0F766E")
PDF_CORAL = colors.HexColor("#D7654D")
PDF_NAVY = colors.HexColor("#18324A")
PDF_LIGHT = colors.HexColor("#E9F2F1")


@dataclass(frozen=True, slots=True)
class RunExportData:
    """Detached, JSON-compatible snapshot of one persisted run."""

    run: dict[str, Any]
    models: list[dict[str, Any]]
    tasks: list[dict[str, Any]]
    aggregates: list[dict[str, Any]]
    items: list[dict[str, Any]]

    def as_dict(self) -> dict[str, Any]:
        """Return the complete versioned JSON export document."""

        return {
            "schema_version": 2,
            "exported_at": _isoformat(datetime.now(UTC)),
            "run": self.run,
            "models": self.models,
            "tasks": self.tasks,
            "aggregates": self.aggregates,
            "items": self.items,
            "caveats": {
                "clinical_use": "Research evaluation only; not a clinical tool.",
                "cost": "All cost values are estimates.",
                "tokens": "Token counts are estimates where tokens_estimated is true.",
            },
        }


@dataclass(frozen=True, slots=True)
class ExportFile:
    """Rendered download plus its HTTP metadata."""

    content: bytes
    mimetype: str
    extension: str


def load_run_export(session: Session, run_id: str) -> RunExportData | None:
    """Load and detach all persisted data required for a run export."""

    run = session.get(Run, run_id)
    if run is None:
        return None
    models = list(
        session.scalars(
            select(RunModel).where(RunModel.run_id == run_id).order_by(RunModel.model_id)
        )
    )
    tasks = list(
        session.scalars(select(RunTask).where(RunTask.run_id == run_id).order_by(RunTask.task_id))
    )
    aggregates = list(
        session.scalars(
            select(Aggregate)
            .where(Aggregate.run_id == run_id)
            .order_by(
                Aggregate.model_id,
                Aggregate.task_id,
                Aggregate.condition_id,
                Aggregate.stat_name,
                Aggregate.field,
            )
        )
    )
    results = list(
        session.scalars(
            select(Result)
            .where(Result.run_id == run_id)
            .order_by(Result.model_id, Result.task_id, Result.item_id, Result.repetition)
        )
    )
    return RunExportData(
        run={
            "id": run.id,
            "label": run.label,
            "status": run.status,
            "created_at": _isoformat(run.created_at),
            "started_at": _isoformat(run.started_at),
            "finished_at": _isoformat(run.finished_at),
            "settings": run.settings_json,
            "app_version": run.app_version,
            "litellm_host": run.litellm_host,
            "notes": run.notes,
        },
        models=[
            {"model_id": model.model_id, "pricing_source": model.pricing_source_summary}
            for model in models
        ],
        tasks=[
            {
                "task_id": task.task_id,
                "version": task.task_version,
                "task_file_sha256": task.task_file_sha256,
                "gold_file_sha256": task.gold_file_sha256,
                "prompt_sha256": task.prompt_sha256,
            }
            for task in tasks
        ],
        aggregates=[
            {
                "run_id": aggregate.run_id,
                "model_id": aggregate.model_id,
                "task_id": aggregate.task_id,
                "condition_id": aggregate.condition_id,
                "stat_name": aggregate.stat_name,
                "field": aggregate.field,
                "value": aggregate.value,
                "n": aggregate.n,
            }
            for aggregate in aggregates
        ],
        items=[_serialize_result(result) for result in results],
    )


def render_export(data: RunExportData, export_format: str) -> ExportFile:
    """Render *data* to a supported format."""

    if export_format == "json":
        return ExportFile(
            json.dumps(data.as_dict(), ensure_ascii=False, indent=2).encode("utf-8"),
            "application/json",
            "json",
        )
    if export_format == "csv":
        return ExportFile(_csv_zip(data), "application/zip", "csv.zip")
    if export_format == "xlsx":
        return ExportFile(
            _xlsx(data),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "xlsx",
        )
    if export_format == "pdf":
        return ExportFile(_pdf(data), "application/pdf", "pdf")
    raise ValueError(f"Unsupported export format: {export_format}")


def export_filename(data: RunExportData, extension: str) -> str:
    """Return a filesystem-safe attachment name that cannot expose user content."""

    short_id = re.sub(r"[^A-Za-z0-9-]", "", data.run["id"])[:36] or "run"
    return f"aphasia-eval-{short_id}.{extension}"


def _serialize_result(result: Result) -> dict[str, Any]:
    experiment = result_experiment(result.comparison_json, result.repetition)
    return {
        "result_id": result.id,
        "run_id": result.run_id,
        "model_id": result.model_id,
        "task_id": result.task_id,
        "item_id": result.item_id,
        "condition_id": experiment["condition_id"],
        "condition_json": experiment,
        "repetition": experiment["repetition"],
        "input_json": result.input_json,
        "gold_json": result.gold_json,
        "prompt_rendered": result.prompt_rendered,
        "raw_output": result.raw_output,
        "parsed_output_json": result.parsed_output_json,
        "derived_json": result.derived_json,
        "comparison_json": result.comparison_json,
        "matched": result.matched,
        "input_tokens": result.input_tokens,
        "output_tokens": result.output_tokens,
        "reasoning_tokens": result.reasoning_tokens,
        "tokens_estimated": result.tokens_estimated,
        "estimated_cost_usd": result.cost_usd,
        "cost_source": result.cost_source,
        "latency_ms": result.latency_ms,
        "total_elapsed_ms": result.total_elapsed_ms,
        "attempts": result.attempts,
        "status": result.status,
        "error_message": result.error_message,
        "created_at": _isoformat(result.created_at),
    }


def _csv_zip(data: RunExportData) -> bytes:
    output = BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("items.csv", _csv_text(CSV_ITEM_COLUMNS, data.items))
        columns, rows = _wide_aggregate_rows(data)
        archive.writestr("aggregates.csv", _csv_text(columns, rows))
    return output.getvalue()


def _wide_aggregate_rows(data: RunExportData) -> tuple[tuple[str, ...], list[dict[str, Any]]]:
    metric_names = sorted({_aggregate_column(entry) for entry in data.aggregates})
    columns = (*CSV_AGGREGATE_ID_COLUMNS, *metric_names, *(f"{name}__n" for name in metric_names))
    grouped: dict[tuple[str, str, str], dict[str, Any]] = {}
    for aggregate in data.aggregates:
        key = (
            aggregate["model_id"],
            aggregate["task_id"],
            aggregate["condition_id"],
        )
        row = grouped.setdefault(
            key,
            {
                "run_id": aggregate["run_id"],
                "model_id": aggregate["model_id"],
                "task_id": aggregate["task_id"],
                "condition_id": aggregate["condition_id"],
            },
        )
        metric = _aggregate_column(aggregate)
        row[metric] = aggregate["value"]
        row[f"{metric}__n"] = aggregate["n"]
    return tuple(columns), [grouped[key] for key in sorted(grouped)]


def _aggregate_column(aggregate: dict[str, Any]) -> str:
    field = aggregate.get("field")
    return f"{aggregate['stat_name']}.{field}" if field else aggregate["stat_name"]


def _csv_text(columns: tuple[str, ...], rows: list[dict[str, Any]]) -> str:
    stream = StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=columns, extrasaction="ignore", lineterminator="\n")
    writer.writeheader()
    for row in rows:
        writer.writerow({column: _tabular_value(row.get(column)) for column in columns})
    return stream.getvalue()


def _tabular_value(value: Any) -> Any:
    if isinstance(value, dict | list):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
        return "'" + value
    return value


def _xlsx(data: RunExportData) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    _write_summary_sheet(summary, data)
    _write_efficiency_sheet(workbook.create_sheet("Efficiency"), data)
    _write_items_sheet(workbook.create_sheet("Items"), data)
    if isinstance(data.run["settings"].get("ablation"), dict):
        _write_ablation_sheet(workbook.create_sheet("Ablation"), data)
    _write_run_info_sheet(workbook.create_sheet("Run Info"), data)
    for sheet in workbook.worksheets:
        _finish_sheet(sheet)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _write_summary_sheet(sheet, data: RunExportData) -> None:
    sheet.append(["aphasia-llm-eval run summary"])
    sheet.append(["Run", _excel_value(data.run["label"] or data.run["id"])])
    sheet.append(["Status", data.run["status"]])
    sheet.append(["Clinical-use caveat", "Research evaluation only; not a clinical tool."])
    sheet.append(["Cost caveat", "All cost values are estimated."])
    sheet.append([])
    sheet.append(["Model", "Task", "Condition", "Statistic", "Field", "Value", "N"])
    for aggregate in data.aggregates:
        sheet.append(
            [
                _excel_value(aggregate["model_id"]),
                _excel_value(aggregate["task_id"]),
                _excel_value(aggregate["condition_id"]),
                _excel_value(aggregate["stat_name"]),
                _excel_value(aggregate["field"]),
                aggregate["value"],
                aggregate["n"],
            ]
        )
    sheet.freeze_panes = "A8"
    sheet.auto_filter.ref = f"A7:G{max(7, sheet.max_row)}"


def _write_efficiency_sheet(sheet, data: RunExportData) -> None:
    columns = (
        "Model",
        "Task",
        "Condition",
        "Input tokens",
        "Output tokens",
        "Reasoning tokens",
        "Estimated cost (USD)",
        "Cost source",
        "Total time (ms)",
        "Mean latency (ms)",
        "P95 latency (ms)",
        "Any token counts estimated",
    )
    sheet.append(columns)
    aggregate_lookup = _condition_aggregate_lookup(data)
    for model in data.models:
        model_id = model["model_id"]
        for task in data.tasks:
            task_id = task["task_id"]
            condition_ids = sorted(
                {
                    entry["condition_id"]
                    for entry in data.aggregates
                    if entry["model_id"] == model_id and entry["task_id"] == task_id
                },
                key=lambda value: value != "baseline",
            )
            for condition_id in condition_ids:
                estimated = any(
                    item["model_id"] == model_id
                    and item["task_id"] == task_id
                    and item["condition_id"] == condition_id
                    and item["tokens_estimated"]
                    for item in data.items
                )
                key = (model_id, task_id, condition_id)
                sheet.append(
                    [
                        _excel_value(model_id),
                        _excel_value(task_id),
                        _excel_value(condition_id),
                        aggregate_lookup.get((*key, "input_tokens", None)),
                        aggregate_lookup.get((*key, "output_tokens", None)),
                        aggregate_lookup.get((*key, "reasoning_tokens", None)),
                        aggregate_lookup.get((*key, "estimated_cost_usd", None)),
                        model["pricing_source"] or "unknown",
                        aggregate_lookup.get((*key, "total_elapsed_ms", None)),
                        aggregate_lookup.get((*key, "mean_latency_ms", None)),
                        aggregate_lookup.get((*key, "p95_latency_ms", None)),
                        estimated,
                    ]
                )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:L{max(1, sheet.max_row)}"


def _write_ablation_sheet(sheet, data: RunExportData) -> None:
    """Write a condition-oriented table suitable for independent analysis."""

    columns, rows = _wide_aggregate_rows(data)
    sheet.append([_humanize(column) for column in columns])
    for row in rows:
        sheet.append([_excel_value(row.get(column)) for column in columns])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(1, sheet.max_row)}"


def _write_items_sheet(sheet, data: RunExportData) -> None:
    sheet.append(CSV_ITEM_COLUMNS)
    for item in data.items:
        sheet.append([_excel_value(item.get(column)) for column in CSV_ITEM_COLUMNS])
    sheet.freeze_panes = "A2"
    last_column = get_column_letter(len(CSV_ITEM_COLUMNS))
    sheet.auto_filter.ref = f"A1:{last_column}{max(1, sheet.max_row)}"


def _write_run_info_sheet(sheet, data: RunExportData) -> None:
    sheet.append(["Field", "Value"])
    for key, value in data.run.items():
        sheet.append([key, _excel_value(value)])
    sheet.append([])
    sheet.append(["Models"])
    sheet.append(["Model", "Pricing source"])
    for model in data.models:
        sheet.append([_excel_value(model["model_id"]), model["pricing_source"] or "unknown"])
    sheet.append([])
    sheet.append(["Tasks"])
    sheet.append(["Task", "Version", "Task SHA-256", "Prompt SHA-256", "Gold SHA-256"])
    for task in data.tasks:
        sheet.append(
            [
                _excel_value(task["task_id"]),
                _excel_value(task["version"]),
                task["task_file_sha256"],
                task["prompt_sha256"],
                task["gold_file_sha256"],
            ]
        )


def _excel_value(value: Any) -> Any:
    value = _tabular_value(value)
    if isinstance(value, str) and len(value) > 32767:
        return value[:32710] + "… [truncated for XLSX; use JSON export for full value]"
    return value


def _finish_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="0F766E")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column in range(1, sheet.max_column + 1):
        values = (
            str(sheet.cell(row=row, column=column).value or "")
            for row in range(1, sheet.max_row + 1)
        )
        width = min(60, max(12, max((len(value) for value in values), default=12) + 2))
        sheet.column_dimensions[get_column_letter(column)].width = width


def _pdf(data: RunExportData) -> bytes:
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(letter),
        rightMargin=0.45 * inch,
        leftMargin=0.45 * inch,
        topMargin=0.55 * inch,
        bottomMargin=0.55 * inch,
        title=f"aphasia-llm-eval run {data.run['id']}",
        author="aphasia-llm-eval",
        subject="Research evaluation export",
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="CenteredNote", parent=styles["BodyText"], alignment=TA_CENTER))
    story: list[Any] = [
        Paragraph("aphasia-llm-eval", styles["Title"]),
        Paragraph(_safe_text(data.run["label"] or f"Run {data.run['id']}"), styles["Heading2"]),
        Paragraph(
            "<b>Research evaluation only.</b> Not a clinical tool. Do not use this report "
            "for diagnosis or treatment decisions.",
            styles["CenteredNote"],
        ),
        Paragraph("All displayed cost values are estimated.", styles["CenteredNote"]),
        Spacer(1, 12),
        Paragraph("Run information", styles["Heading2"]),
        _pdf_table(
            [["Field", "Value"]]
            + [
                [_humanize(key), _safe_text(_display_value(value))]
                for key, value in data.run.items()
            ],
            [1.7 * inch, 7.7 * inch],
        ),
        Spacer(1, 14),
        Paragraph("Overview matrix", styles["Heading2"]),
        _overview_table(data),
        Spacer(1, 14),
        Paragraph("Efficiency", styles["Heading2"]),
        _efficiency_table(data),
        Spacer(1, 10),
    ]
    cost_chart = _cost_chart(data)
    if cost_chart is not None:
        story.extend([Paragraph("Estimated cost by model", styles["Heading3"]), cost_chart])
    story.append(PageBreak())
    aggregate_lookup = _aggregate_lookup(data)
    for task in data.tasks:
        task_id = task["task_id"]
        story.append(Paragraph(f"Task summary: {_safe_text(task_id)}", styles["Heading2"]))
        task_rows = [["Model", "Condition", "Statistic", "Field", "Value", "N"]]
        for aggregate in data.aggregates:
            if aggregate["task_id"] == task_id:
                task_rows.append(
                    [
                        _safe_text(aggregate["model_id"]),
                        _safe_text(aggregate["condition_id"]),
                        _safe_text(aggregate["stat_name"]),
                        _safe_text(aggregate["field"] or "—"),
                        _format_number(aggregate["value"]),
                        str(aggregate["n"]),
                    ]
                )
        story.append(
            _pdf_table(
                task_rows,
                [1.6 * inch, 1.25 * inch, 1.7 * inch, 1.25 * inch, 1 * inch, 0.5 * inch],
            )
        )
        chart = _headline_chart(data, task_id, aggregate_lookup)
        if chart is not None:
            story.extend([Spacer(1, 8), chart])
        story.append(Spacer(1, 18))
    story.extend(
        [
            PageBreak(),
            Paragraph("Methods and limitations", styles["Heading2"]),
            Paragraph(
                "Each model response was compared with the configured task's gold answer using "
                "that task's comparator. Headline statistics are reported separately by task; "
                "this report does not calculate a cross-task composite score.",
                styles["BodyText"],
            ),
            Spacer(1, 8),
            Paragraph(
                "Agreement with a gold file does not establish clinical validity. Results depend "
                "on prompt wording, task definitions, gold provenance, model versions, and gateway "
                "behavior. Token counts marked as estimated use a fallback tokenizer. Costs are "
                "estimates and may differ from provider billing.",
                styles["BodyText"],
            ),
            Spacer(1, 12),
            Paragraph("Task and gold provenance", styles["Heading3"]),
            _pdf_table(
                [["Task", "Version", "Task SHA-256", "Prompt SHA-256", "Gold SHA-256"]]
                + [
                    [
                        _safe_text(task["task_id"]),
                        _safe_text(task["version"]),
                        task["task_file_sha256"],
                        task["prompt_sha256"],
                        task["gold_file_sha256"],
                    ]
                    for task in data.tasks
                ],
                [1.25 * inch, 0.65 * inch, 2.5 * inch, 2.5 * inch, 2.5 * inch],
                font_size=6,
            ),
        ]
    )
    document.build(story, onFirstPage=_pdf_footer, onLaterPages=_pdf_footer)
    return output.getvalue()


def _overview_table(data: RunExportData) -> Table:
    aggregate_lookup = _aggregate_lookup(data)
    rows = [["Model", *[task["task_id"] for task in data.tasks]]]
    for model in data.models:
        model_id = model["model_id"]
        rows.append(
            [
                _safe_text(model_id),
                *[
                    _format_number(
                        aggregate_lookup.get(
                            (
                                model_id,
                                task["task_id"],
                                "headline",
                                _headline_field(data, model_id, task["task_id"]),
                            )
                        )
                    )
                    for task in data.tasks
                ],
            ]
        )
    widths = [2.4 * inch] + [min(2.2 * inch, 7 * inch / max(1, len(data.tasks)))] * len(data.tasks)
    return _pdf_table(rows, widths)


def _efficiency_table(data: RunExportData) -> Table:
    lookup = _aggregate_lookup(data)
    rows = [
        [
            "Model",
            "Task",
            "Input tokens",
            "Output tokens",
            "Estimated cost",
            "Source",
            "Mean ms",
            "P95 ms",
        ]
    ]
    for model in data.models:
        model_id = model["model_id"]
        for task in data.tasks:
            task_id = task["task_id"]
            estimated = any(
                item["model_id"] == model_id
                and item["task_id"] == task_id
                and item["tokens_estimated"]
                for item in data.items
            )
            token_suffix = " est." if estimated else ""
            rows.append(
                [
                    _safe_text(model_id),
                    _safe_text(task_id),
                    _format_number(lookup.get((model_id, task_id, "input_tokens", None)))
                    + token_suffix,
                    _format_number(lookup.get((model_id, task_id, "output_tokens", None)))
                    + token_suffix,
                    _format_currency(lookup.get((model_id, task_id, "estimated_cost_usd", None))),
                    _safe_text(model["pricing_source"] or "unknown"),
                    _format_number(lookup.get((model_id, task_id, "mean_latency_ms", None))),
                    _format_number(lookup.get((model_id, task_id, "p95_latency_ms", None))),
                ]
            )
    return _pdf_table(
        rows,
        [
            1.8 * inch,
            1.3 * inch,
            1 * inch,
            1 * inch,
            1.1 * inch,
            0.9 * inch,
            0.85 * inch,
            0.85 * inch,
        ],
        font_size=7,
    )


def _pdf_table(rows: list[list[Any]], widths: list[float], font_size: int = 8) -> Table:
    cell_style = ParagraphStyle(
        name=f"ExportCell{font_size}",
        fontName="Helvetica",
        fontSize=font_size,
        leading=font_size + 2,
        wordWrap="CJK",
    )
    wrapped_rows = [
        rows[0],
        *[
            [Paragraph(value, cell_style) if isinstance(value, str) else value for value in row]
            for row in rows[1:]
        ],
    ]
    table = Table(wrapped_rows, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), PDF_NAVY),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (0, 1), (0, -1), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), font_size),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#B6C4C2")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, PDF_LIGHT]),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return table


def _cost_chart(data: RunExportData) -> Drawing | None:
    lookup = _aggregate_lookup(data)
    labels: list[str] = []
    values: list[float] = []
    for model in data.models:
        model_id = model["model_id"]
        cost_values = [
            lookup.get((model_id, task["task_id"], "estimated_cost_usd", None))
            for task in data.tasks
        ]
        known = [float(value) for value in cost_values if value is not None]
        if known:
            labels.append(model_id)
            values.append(sum(known))
    return _bar_chart(labels, values, PDF_CORAL)


def _headline_chart(
    data: RunExportData,
    task_id: str,
    lookup: dict[tuple[str, str, str, str | None], float | None],
) -> Drawing | None:
    labels: list[str] = []
    values: list[float] = []
    for model in data.models:
        model_id = model["model_id"]
        field = _headline_field(data, model_id, task_id)
        value = lookup.get((model_id, task_id, "headline", field))
        if value is not None:
            labels.append(model_id)
            values.append(float(value))
    return _bar_chart(labels, values, PDF_BLUE)


def _bar_chart(labels: list[str], values: list[float], color) -> Drawing | None:
    if not values:
        return None
    drawing = Drawing(680, 220)
    chart = VerticalBarChart()
    chart.x = 55
    chart.y = 45
    chart.height = 145
    chart.width = 590
    chart.data = [values]
    chart.categoryAxis.categoryNames = [label[:24] for label in labels]
    chart.categoryAxis.labels.angle = 25
    chart.categoryAxis.labels.fontSize = 6
    chart.valueAxis.valueMin = min(0, min(values))
    maximum = max(values)
    chart.valueAxis.valueMax = maximum if maximum > 0 else 1
    chart.valueAxis.labelTextFormat = "%0.3g"
    chart.bars[0].fillColor = color
    chart.bars[0].strokeColor = color
    drawing.add(chart)
    return drawing


def _aggregate_lookup(
    data: RunExportData,
) -> dict[tuple[str, str, str, str | None], float | None]:
    return {
        (entry["model_id"], entry["task_id"], entry["stat_name"], entry["field"]): entry["value"]
        for entry in data.aggregates
        if entry["condition_id"] == "baseline"
    }


def _condition_aggregate_lookup(
    data: RunExportData,
) -> dict[tuple[str, str, str, str, str | None], float | None]:
    return {
        (
            entry["model_id"],
            entry["task_id"],
            entry["condition_id"],
            entry["stat_name"],
            entry["field"],
        ): entry["value"]
        for entry in data.aggregates
    }


def _headline_field(data: RunExportData, model_id: str, task_id: str) -> str | None:
    entry = next(
        (
            aggregate
            for aggregate in data.aggregates
            if aggregate["model_id"] == model_id
            and aggregate["task_id"] == task_id
            and aggregate["stat_name"] == "headline"
        ),
        None,
    )
    return entry["field"] if entry else None


def _pdf_footer(canvas, document) -> None:
    canvas.saveState()
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(colors.HexColor("#5D6D6A"))
    canvas.drawString(
        document.leftMargin, 0.3 * inch, "aphasia-llm-eval · research evaluation only"
    )
    canvas.drawRightString(
        landscape(letter)[0] - document.rightMargin,
        0.3 * inch,
        f"Page {document.page}",
    )
    canvas.restoreState()


def _safe_text(value: Any) -> str:
    return escape(str(value), {'"': "&quot;", "'": "&#39;"})


def _display_value(value: Any) -> str:
    if isinstance(value, dict | list):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return "n/a" if value is None else str(value)


def _humanize(value: str) -> str:
    return value.replace("_", " ").capitalize()


def _format_number(value: float | int | None) -> str:
    if value is None or not math.isfinite(float(value)):
        return "n/a"
    return f"{float(value):,.3f}".rstrip("0").rstrip(".")


def _format_currency(value: float | int | None) -> str:
    return "n/a" if value is None else f"${float(value):,.6f} est."


def _isoformat(value: datetime | None) -> str | None:
    if value is None:
        return None
    if value.tzinfo is None:
        value = value.replace(tzinfo=UTC)
    return value.astimezone(UTC).isoformat().replace("+00:00", "Z")
